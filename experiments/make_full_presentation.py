"""Build the full-study (11 models) results deck from scored results and the human audit.

Every number on the slides is computed from:
    results/scored/test_final.csv
    results/analysis/test_final/{summary, effort_effects}.csv
    results/audit/distractor_audit_final.csv
Output: results/presentation/webshop_full_study_results.pptx (+ figures_full/)

Usage (repo root; pandas, matplotlib, python-pptx):
    python experiments/make_full_presentation.py
"""
import os
import sys
from os.path import abspath, dirname, join

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

sys.path.insert(0, dirname(abspath(__file__)))
from make_presentation import NAVY, Deck, pct, wilson  # noqa: E402

REPO_DIR = dirname(dirname(abspath(__file__)))
RESULTS = join(REPO_DIR, 'results')
OUT_DIR = join(RESULTS, 'presentation')
FIG_DIR = join(OUT_DIR, 'figures_full')
SPLIT = 'test_final'

MODELS = ['gpt-6-astra', 'gpt-6.1-sol', 'gpt-6-luna', 'claude-fable-5.1', 'claude-opus-5.5',
          'gemini-4-argon', 'gemini-3.1-pro', 'deepseek-v4-pro', 'qwen3.8-max', 'kimi-k3', 'glm-5v-turbo']
NAMES = {'gpt-6-astra': 'GPT-6 Astra', 'gpt-6.1-sol': 'GPT-6.1 Sol', 'gpt-6-luna': 'GPT-6 Luna',
         'claude-fable-5.1': 'Claude Fable 5.1', 'claude-opus-5.5': 'Claude Opus 5.5',
         'gemini-4-argon': 'Gemini 4 Argon', 'gemini-3.1-pro': 'Gemini 3.1 Pro',
         'deepseek-v4-pro': 'DeepSeek V4 Pro', 'qwen3.8-max': 'Qwen 3.8 Max', 'kimi-k3': 'Kimi K3',
         'glm-5v-turbo': 'GLM-5V Turbo'}
COLORS = {'gpt-6-astra': '#1b7837', 'gpt-6.1-sol': '#5aae61', 'gpt-6-luna': '#a6dba0',
          'claude-fable-5.1': '#08306b', 'claude-opus-5.5': '#4292c6',
          'gemini-4-argon': '#b2182b', 'gemini-3.1-pro': '#f4a582',
          'deepseek-v4-pro': '#762a83', 'qwen3.8-max': '#e08214', 'kimi-k3': '#8c510a',
          'glm-5v-turbo': '#7f7f7f'}
EFFORTS = ['low', 'medium', 'high']
EFFORT_X = {'low': 0, 'medium': 1, 'high': 2, 'default': 1}


def load():
    d = {
        'scored': pd.read_csv(join(RESULTS, 'scored', f'{SPLIT}.csv')),
        'summary': pd.read_csv(join(RESULTS, 'analysis', SPLIT, 'summary.csv')),
        'effects': pd.read_csv(join(RESULTS, 'analysis', SPLIT, 'effort_effects.csv')),
        'audit': pd.read_csv(join(RESULTS, 'audit', 'distractor_audit_final.csv'), encoding='utf-8-sig'),
    }
    v = d['scored'][d['scored']['valid'].astype(bool)].copy()
    v['success'] = v['success'].astype(bool)
    v['success_lenient'] = v['success_lenient'].astype(bool)
    d['valid'] = v
    return d


def poisson_binomial(ps):
    dist = np.zeros(len(ps) + 1)
    dist[0] = 1.0
    for p in ps:
        dist[1:] = dist[1:] * (1 - p) + dist[:-1] * p
        dist[0] *= (1 - p)
    return dist


def compute_stats(d, rng):
    v, summary, effects, audit = d['valid'], d['summary'], d['effects'], d['audit']
    s = {'n_tasks': d['scored']['task_id'].nunique(), 'n_responses': len(d['scored']), 'n_valid': len(v),
         'n_conditions': d['scored'].groupby(['model_key', 'effort']).ngroups}
    s['summary'] = summary.set_index(['model_key', 'effort'])
    s['hl'] = effects[effects['contrast'] == 'high - low'].set_index(['model_key', 'metric'])

    verdicts = audit['violates_instruction'].str.strip().str.lower()
    s['audit_n'], s['audit_false'] = len(audit), int((verdicts == 'n').sum())
    s['audit_ci'] = wilson(s['audit_false'], s['audit_n'])
    gap = (v.groupby(['model_key', 'effort'])['success_lenient'].mean()
           - v.groupby(['model_key', 'effort'])['success'].mean())
    s['lenient_gap'] = (gap.min(), gap.max())

    s['conf_rw'] = v.groupby(['model_key', 'success'])['confidence'].mean().unstack()
    s['share10'] = v.assign(c10=v['confidence'] == 10).groupby('model_key')['c10'].mean()
    s['tokens'] = v.groupby(['model_key', 'effort'])['reasoning_tokens'].mean()

    # Concentration: each model at its lowest available effort (low, or default for GLM).
    first = {m: ('low' if (m, 'low') in s['summary'].index else 'default') for m in MODELS}
    base = v[[first[m] == e for m, e in zip(v['model_key'], v['effort'])]]
    wrong = base.pivot_table(index='task_id', columns='model_key', values='success', aggfunc='first')
    wrong = ~wrong.astype(bool)
    k_obs = wrong.sum(axis=1).astype(int).value_counts().reindex(range(len(MODELS) + 1), fill_value=0)
    err_rates = wrong.mean()
    k_exp = poisson_binomial(err_rates.values) * len(wrong)
    s['conc_k_obs'], s['conc_k_exp'] = k_obs, k_exp
    s['conc_threshold'] = 6
    s['conc_obs_maj'] = int(k_obs[s['conc_threshold']:].sum())
    s['conc_exp_maj'] = float(k_exp[s['conc_threshold']:].sum())
    s['conc_obs_zero'], s['conc_exp_zero'] = int(k_obs[0]), float(k_exp[0])
    picks = base[~base['success']].groupby('task_id')['chosen_asin'].apply(list)
    picks = picks[picks.apply(len) >= 2]
    modal = picks.apply(lambda xs: pd.Series(xs).value_counts().iloc[0] / len(xs)).mean()
    sim = np.mean([np.mean([np.bincount(rng.integers(0, 9, len(xs)), minlength=9).max() / len(xs)
                            for xs in picks]) for _ in range(200)])
    s['conc_modal'], s['conc_modal_random'], s['conc_n_multi'] = modal, sim, len(picks)

    # Position bias (all valid answers).
    s['pos_share'] = v['chosen_position'].value_counts(normalize=True).sort_index()
    s['succ_by_target_pos'] = v.groupby('target_position')['success'].mean()
    s['succ_front'] = v[v['target_position'] <= 5]['success'].mean()
    s['succ_back'] = v[v['target_position'] > 5]['success'].mean()
    return s


def model_levels(summary, m):
    return [e for e in EFFORTS + ['default'] if (m, e) in summary.index]


def fig_effort(s, path):
    sm = s['summary']
    fig, axes = plt.subplots(1, 3, figsize=(15.5, 4.6))
    for m in MODELS:
        levels = model_levels(sm, m)
        g = sm.loc[[(m, e) for e in levels]]
        x = [EFFORT_X[e] for e in levels]
        default = levels == ['default']
        style = dict(color=COLORS[m], marker='D' if default else 'o', linestyle='none' if default else '-',
                     linewidth=2, markersize=6, label=NAMES[m] + (' (no effort control)' if default else ''))
        axes[0].plot(x, 100 * g['success'], **style)
        axes[1].plot(x, g['mean_confidence'], **style)
        axes[2].plot(x, g['mean_reasoning_tokens'], **style)
    for ax, title in zip(axes, ['Correctness: success rate (%)', 'Stated confidence (0-10)',
                                'Reasoning tokens used (mean)']):
        ax.set_xticks(range(3))
        ax.set_xticklabels(EFFORTS)
        ax.set_xlabel('reasoning effort setting')
        ax.set_title(title, fontsize=12)
        ax.spines[['top', 'right']].set_visible(False)
    axes[0].set_ylim(65, 95)
    axes[1].set_ylim(7, 10)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc='center right', frameon=False, fontsize=9)
    fig.tight_layout(rect=(0, 0, 0.83, 1))
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_effects(s, path):
    hl = s['hl']
    models = [m for m in MODELS if (m, 'success') in hl.index]
    y = np.arange(len(models))[::-1]
    fig, ax = plt.subplots(figsize=(8.2, 5.4))
    for yi, m in zip(y, models):
        r = hl.loc[(m, 'success')]
        ax.errorbar(100 * r['diff'], yi + 0.15, xerr=[[100 * (r['diff'] - r['ci_low'])], [100 * (r['ci_high'] - r['diff'])]],
                    fmt='o', color='#2a9d8f', capsize=3)
        c = hl.loc[(m, 'confidence')]
        ax.errorbar(10 * c['diff'], yi - 0.15, xerr=[[10 * (c['diff'] - c['ci_low'])], [10 * (c['ci_high'] - c['diff'])]],
                    fmt='s', color='#e76f51', capsize=3)
    ax.axvline(0, color='grey', linewidth=1)
    ax.set_yticks(y)
    ax.set_yticklabels([NAMES[m] for m in models])
    ax.set_xlabel('change from low to high effort (percentage points)')
    ax.plot([], [], 'o', color='#2a9d8f', label='success rate')
    ax.plot([], [], 's', color='#e76f51', label='stated confidence (0-10 scale x 10)')
    ax.legend(frameon=False, loc='lower right', fontsize=9)
    ax.set_title('Paired effect of raising effort (95% bootstrap intervals)', fontsize=11)
    ax.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_calibration(s, path):
    sm = s['summary']
    fig, ax = plt.subplots(figsize=(7.4, 6.0))
    ax.plot([60, 100], [60, 100], color='grey', linestyle='--', linewidth=1)
    ax.text(73.5, 75.5, 'perfectly calibrated', color='grey', fontsize=9, rotation=40)
    sizes = {'low': 40, 'medium': 90, 'high': 160, 'default': 90}
    for m in MODELS:
        levels = model_levels(sm, m)
        g = sm.loc[[(m, e) for e in levels]]
        ax.plot(10 * g['mean_confidence'], 100 * g['success'], color=COLORS[m], linewidth=1, alpha=0.7)
        for e in levels:
            ax.scatter(10 * g.loc[(m, e), 'mean_confidence'], 100 * g.loc[(m, e), 'success'], s=sizes[e],
                       color=COLORS[m], edgecolor='white', marker='D' if e == 'default' else 'o', zorder=3)
        last = g.iloc[-1]
        offset = {'gpt-6-luna': (-58, -12), 'kimi-k3': (6, -10), 'gpt-6.1-sol': (6, -8)}.get(
            m, (6, 4) if m.startswith('claude') else (6, -2))
        ax.annotate(NAMES[m], (10 * last['mean_confidence'], 100 * last['success']), textcoords='offset points',
                    xytext=offset, fontsize=8, color=COLORS[m])
    ax.set_xlim(75, 100)
    ax.set_ylim(70, 92)
    ax.set_xlabel('mean stated confidence (%)')
    ax.set_ylabel('actual success rate (%)')
    ax.set_title('Below the diagonal = overconfident (dot size: low / medium / high effort)', fontsize=10)
    ax.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_right_wrong(s, path):
    crw = s['conf_rw']
    order = sorted(MODELS, key=lambda m: crw.loc[m, True] - crw.loc[m, False], reverse=True)
    x = np.arange(len(order))
    fig, ax = plt.subplots(figsize=(11.5, 4.4))
    right = [crw.loc[m, True] for m in order]
    wrong = [crw.loc[m, False] for m in order]
    ax.bar(x - 0.2, right, 0.4, label='when correct', color='#8ab17d')
    ax.bar(x + 0.2, wrong, 0.4, label='when wrong', color='#e76f51')
    for i in range(len(order)):
        ax.text(x[i], max(right[i], wrong[i]) + 0.06, f'{right[i] - wrong[i]:+.2f}', ha='center', fontsize=9)
    ax.set_xticks(x)
    ax.set_xticklabels([NAMES[m] for m in order], rotation=25, ha='right')
    ax.set_ylim(6.5, 10.5)
    ax.set_ylabel('mean stated confidence (0-10)')
    ax.set_title('Confidence when correct vs. wrong (number above bars = drop when wrong)', fontsize=11)
    ax.legend(frameon=False, loc='upper right', ncol=2)
    ax.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_concentration(s, path):
    k = np.arange(len(MODELS) + 1)
    fig, ax = plt.subplots(figsize=(8.2, 4.6))
    ax.bar(k - 0.2, s['conc_k_obs'].values, 0.4, color='#1d3557', label='observed')
    ax.bar(k + 0.2, s['conc_k_exp'], 0.4, color='#a8dadc', label='expected if errors were independent')
    ax.set_xticks(k)
    ax.set_xlabel(f'number of models (of {len(MODELS)}) that chose wrong on a task')
    ax.set_ylabel(f'number of tasks (of {s["n_tasks"]})')
    ax.set_title('Models fail on the same tasks (each model at its lowest effort setting)', fontsize=11)
    ax.legend(frameon=False)
    ax.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_position(s, path):
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    pos = s['pos_share']
    axes[0].bar(pos.index, 100 * pos.values, color='#457b9d')
    axes[0].axhline(10, color='grey', linestyle='--', linewidth=1)
    axes[0].set_xlabel('position of product in the list')
    axes[0].set_ylabel('share of all choices (%)')
    axes[0].set_title('Which position models choose (dashed: no bias)', fontsize=11)
    tp = s['succ_by_target_pos']
    axes[1].bar(tp.index, 100 * tp.values, color='#2a9d8f')
    axes[1].set_xlabel('position of the correct product')
    axes[1].set_ylabel('success rate (%)')
    axes[1].set_ylim(50, 100)
    axes[1].set_title('Success by where the correct product appears', fontsize=11)
    for ax in axes:
        ax.set_xticks(range(1, 11))
        ax.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def build(s, figs, out_path):
    sm, hl = s['summary'], s['hl']
    deck = Deck()
    white, pale = RGBColor(255, 255, 255), RGBColor(0xd8, 0xe2, 0xec)
    n_models = len(MODELS)

    def rng_of(m, col, fmt, scale=1.0, suffix=''):
        vals = scale * sm.loc[m][col].dropna()
        lo, hi = format(vals.min(), fmt), format(vals.max(), fmt)
        return f'{lo}{suffix}' if lo == hi else f'{lo}-{hi}{suffix}'

    def eff(m, metric, scale, decimals):
        r = hl.loc[(m, metric)]
        vals = [round(scale * r[k], decimals) + 0.0 for k in ['diff', 'ci_low', 'ci_high']]
        return f'{vals[0]:+.{decimals}f} [{vals[1]:+.{decimals}f}, {vals[2]:+.{decimals}f}]'

    def sig(m, metric):
        r = hl.loc[(m, metric)]
        return r['ci_low'] > 0 or r['ci_high'] < 0

    effort_models = [m for m in MODELS if (m, 'success') in hl.index]
    conf_up = [m for m in effort_models if hl.loc[(m, 'confidence'), 'ci_low'] > 0]
    succ_up = [m for m in effort_models if hl.loc[(m, 'success'), 'ci_low'] > 0]
    claude = ['claude-fable-5.1', 'claude-opus-5.5']
    others = [m for m in MODELS if m not in claude]
    over_others = [100 * sm.loc[m]['overconfidence'].mean() for m in others]
    auroc_others = [sm.loc[m]['auroc'].mean() for m in others]
    low_reasoners = [m for m in MODELS if s['tokens'].loc[m].max() < 100]

    # 1. Title
    sl = deck.prs.slides.add_slide(deck.blank)
    bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.prs.slide_width, deck.prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    deck.text(sl, 'Reasoning Effort, Confidence, and Correctness in Frontier AI Shopping Agents',
              Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.8), size=36, bold=True, color=white)
    deck.text(sl, 'Full WebShop study: results', Inches(0.9), Inches(4.0), Inches(11.5), Inches(0.8),
              size=22, color=pale)
    deck.text(sl, f'{s["n_tasks"]} test instructions  |  {n_models} models  |  {s["n_conditions"]} model x effort '
                  f'conditions  |  {s["n_valid"]:,} valid decisions',
              Inches(0.9), Inches(4.9), Inches(11.5), Inches(0.6), size=16, color=pale)

    # 2. Key findings
    sl = deck.slide('Key findings', notes=(
        'The calibration gap is the most robust result. The effort result is more nuanced than in the pilot: '
        'models that actually reason more get slightly more accurate, but their confidence rises in step, '
        'so overconfidence never shrinks.'))
    deck.bullets(sl, [
        f'Calibration gap replicates across {n_models} models: both Claude models are slightly underconfident and '
        f'their confidence tracks correctness; the other {len(others)} are {min(over_others):.0f}-{max(over_others):.0f} '
        f'points overconfident, and their confidence predicts correctness only weakly',
        f'Raising effort raises confidence for {len(conf_up)} of {len(effort_models)} models, but improves accuracy '
        f'reliably for only {len(succ_up)} ({", ".join(NAMES[m] for m in succ_up)}), by at most '
        f'{100 * max(hl.loc[(m, "success"), "diff"] for m in succ_up):.1f} points',
        'Overconfidence is a stable property of each model: it barely changes with effort (only Qwen 3.8 Max, '
        'whose accuracy rises most, moves noticeably closer to calibrated)',
        f'{", ".join(NAMES[m] for m in low_reasoners)} barely reason at any effort setting '
        '(under 100 reasoning tokens on average)',
        'Models fail together: errors concentrate on the same tasks and often the same wrong product',
    ], Inches(0.7), Inches(1.4), Inches(12), Inches(5.6), size=19)

    # 3. Design
    sl = deck.slide('Design', notes=(
        'Same choice-set design as the preliminary study, now with stricter distractors and the full test split.'))
    deck.bullets(sl, [
        f'{s["n_tasks"]} human-written instructions: WebShop\'s full standard test split (4 of 500 dropped because '
        'even their target product cannot earn a perfect score)',
        'Each task: 10 products (text + image), exactly 1 satisfies every requirement',
        '  9 near-miss distractors of the same kind of product (same Amazon query or most specific category), each '
        'failing at least one requirement even under lenient matching',
        'Model returns: chosen product, option values (size, color, ...), confidence 0-10',
        f'{n_models} models via OpenRouter, each at every effort level it supports '
        f'({s["n_conditions"]} conditions):',
        '  low / medium / high for 8 models; low / high for DeepSeek V4 Pro and Kimi K3; '
        'GLM-5V Turbo has no effort control (reasoning on)',
        '  DeepSeek V4 Pro has no image input and sees product text only',
        'Identical prompt, products, order, and output budget for every condition',
    ], Inches(0.7), Inches(1.4), Inches(12), Inches(5.6), size=18)

    # 4. Ground truth
    lo_ci, hi_ci = s['audit_ci']
    sl = deck.slide('Ground-truth validation', notes=(
        'The audit was AI-assisted and reviewed by hand. Distractors that are actually acceptable make the task '
        'harder for every model equally, so comparisons are less affected than absolute success rates.'))
    deck.bullets(sl, [
        'Correctness = WebShop reward of exactly 1.0 (product type, required features, options, price)',
        'Every answer also scored with a lenient matcher (ignores punctuation, hyphens, unit formats, word forms, '
        f'common synonyms): lenient success is {100 * s["lenient_gap"][0]:.1f}-{100 * s["lenient_gap"][1]:.1f} points '
        'higher for every condition, and the effort and calibration results are the same under both',
        f'Human-reviewed audit of {s["audit_n"]} randomly sampled distractors: '
        f'{s["audit_n"] - s["audit_false"]} genuinely violate their instruction',
        f'Residual rate of acceptable distractors: {pct(s["audit_false"] / s["audit_n"])} '
        f'(95% CI {pct(lo_ci)}-{pct(hi_ci)}), mostly paraphrases ("24-hour protection" for "long lasting") '
        'and requirements stated only in a product title',
    ], Inches(0.7), Inches(1.4), Inches(12), Inches(5.6), size=19)

    # 5. Main result
    sl = deck.slide('Reasoning effort: small accuracy gains, rising confidence', notes=(
        'Left: success is nearly flat for most models. Middle: confidence creeps up. Right: only some models '
        'actually use more reasoning when asked.'))
    sl.shapes.add_picture(figs['effort'], Inches(0.2), Inches(1.25), width=Inches(12.9))
    deck.bullets(sl, [
        f'{", ".join(NAMES[m] for m in low_reasoners)} use under 100 reasoning tokens even at high effort',
        f'Largest accuracy gain: {NAMES["qwen3.8-max"]}, almost entirely from low to medium, where its reasoning '
        f'jumps from {s["tokens"].loc[("qwen3.8-max", "low")]:.0f} to {s["tokens"].loc[("qwen3.8-max", "medium")]:.0f} tokens',
    ], Inches(0.6), Inches(5.3), Inches(12.2), Inches(1.6), size=16)

    # 6. Effect sizes
    sl = deck.slide('How much does raising effort change accuracy vs. confidence?', notes=(
        'Both changes on the same percentage-point scale: confidence 0-10 times 10. Where models get more accurate, '
        'confidence rises by a similar amount, which is why overconfidence stays flat.'))
    sl.shapes.add_picture(figs['effects'], Inches(0.3), Inches(1.3), height=Inches(5.6))
    deck.bullets(sl, [
        f'Confidence rises reliably for {len(conf_up)} of {len(effort_models)} models '
        f'(Kimi K3 is the exception: it falls)',
        f'Accuracy rises reliably for {len(succ_up)}: ' + ', '.join(
            f'{NAMES[m]} {100 * hl.loc[(m, "success"), "diff"]:+.1f}' for m in succ_up),
        'For the rest, the accuracy change is not distinguishable from zero',
        'Both changes are small: at most a few percentage points',
    ], Inches(8.6), Inches(1.6), Inches(4.5), Inches(5.2), size=15)

    # 7. Calibration scatter
    sl = deck.slide('Calibration: Claude vs. everyone else', notes=(
        'Overconfidence = mean stated confidence (as a percentage) minus actual success. Effort moves models '
        'roughly parallel to the diagonal, so their distance from it, the overconfidence, stays the same.'))
    sl.shapes.add_picture(figs['calibration'], Inches(0.3), Inches(1.25), height=Inches(5.9))
    deck.bullets(sl, [
        'Both Claude models sit above the diagonal: slightly underconfident',
        f'All other models sit well below it: {min(over_others):.0f}-{max(over_others):.0f} points overconfident',
        'More effort mostly moves models parallel to the diagonal; only Qwen 3.8 Max moves noticeably toward it',
        f'Most accurate model, {NAMES["gemini-4-argon"]} '
        f'({100 * sm.loc["gemini-4-argon"]["success"].max():.0f}%), is still about '
        f'{100 * sm.loc["gemini-4-argon"]["overconfidence"].mean():.0f} points overconfident',
    ], Inches(7.9), Inches(1.6), Inches(5.2), Inches(5.2), size=15)

    # 8. Table
    sl = deck.slide('All models at a glance', notes='Ranges span the effort levels each model supports.')
    rows = [['Model', 'Success', 'Confidence', 'Overconfidence', 'Answers at 10/10', 'AUROC*', 'Reasoning tokens']]
    for m in MODELS:
        g = sm.loc[m]
        rows.append([NAMES[m], rng_of(m, 'success', '.0f', 100, '%'), rng_of(m, 'mean_confidence', '.1f'),
                     f'{100 * g["overconfidence"].mean():+.0f} pts', pct(s['share10'][m]), rng_of(m, 'auroc', '.2f'),
                     rng_of(m, 'mean_reasoning_tokens', '.0f')])
    deck.table(sl, rows, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.0),
               col_widths=[2.3, 1.4, 1.5, 1.8, 1.7, 1.5, 2.1], size=13)
    deck.text(sl, '*AUROC: how well confidence separates right from wrong answers (0.5 = chance, 1.0 = perfect).',
              Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.5), size=11)

    # 9. Right vs wrong
    crw = s['conf_rw']
    drops_others = [crw.loc[m, True] - crw.loc[m, False] for m in others]
    sl = deck.slide("Only Claude's confidence tells you when it is wrong", notes='Pooled across effort levels.')
    sl.shapes.add_picture(figs['right_wrong'], Inches(0.4), Inches(1.3), width=Inches(12.5))
    deck.bullets(sl, [
        f'Claude Fable and Opus drop by {crw.loc["claude-fable-5.1", True] - crw.loc["claude-fable-5.1", False]:.2f} and '
        f'{crw.loc["claude-opus-5.5", True] - crw.loc["claude-opus-5.5", False]:.2f} points when wrong; '
        f'the other models drop by {min(drops_others):.2f}-{max(drops_others):.2f}',
        f'Gemini 3.1 Pro\'s confidence becomes less informative with effort (AUROC '
        f'{sm.loc[("gemini-3.1-pro", "low"), "auroc"]:.2f} at low vs {sm.loc[("gemini-3.1-pro", "high"), "auroc"]:.2f} at high)',
    ], Inches(0.6), Inches(6.0), Inches(12.2), Inches(1.3), size=15)

    # 10. Concentration
    sl = deck.slide('Market concentration: models fail together', notes=(
        'Each model at its lowest effort setting (GLM at its default). Expected counts assume each model errs '
        'independently at its own error rate.'))
    sl.shapes.add_picture(figs['concentration'], Inches(0.3), Inches(1.4), width=Inches(8.0))
    deck.bullets(sl, [
        f'Tasks where at least {s["conc_threshold"]} of {n_models} models chose wrong: {s["conc_obs_maj"]} observed vs '
        f'{s["conc_exp_maj"]:.1f} expected if errors were independent',
        f'Tasks every model got right: {s["conc_obs_zero"]} observed vs {s["conc_exp_zero"]:.0f} expected',
        f'When several models err on the same task, the most common wrong product accounts for '
        f'{pct(s["conc_modal"])} of their wrong answers ({pct(s["conc_modal_random"])} if wrong answers were spread '
        'at random over the 9 distractors)',
    ], Inches(8.6), Inches(1.6), Inches(4.5), Inches(5.2), size=15)

    # 11. Caveats
    position_biased = abs(s['succ_front'] - s['succ_back']) >= 0.03
    position_line = (
        f'Position bias: success is {100 * s["succ_front"]:.0f}% when the correct product is in positions 1-5 vs '
        f'{100 * s["succ_back"]:.0f}% in 6-10 (positions are balanced across conditions; see backup)'
        if position_biased else
        f'No position bias in the full study (success {100 * s["succ_front"]:.0f}% vs {100 * s["succ_back"]:.0f}% '
        'with the correct product in positions 1-5 vs 6-10; see backup)')
    sl = deck.slide('Caveats', notes='Limitations to keep in mind.')
    deck.bullets(sl, [
        f'Weak effort manipulation for {len(low_reasoners)} of {n_models} models: they barely reason at any setting, '
        'so for them the finding is "the effort setting barely changes behavior"',
        f'About {pct(s["audit_false"] / s["audit_n"])} of distractors may be acceptable; this lowers absolute success '
        'for all models, but affects comparisons less',
        position_line,
        'One sample per task and condition; fixed choice sets rather than interactive search',
        'DeepSeek V4 Pro sees text only; no text-only vs image comparison yet for the other models',
    ], Inches(0.7), Inches(1.4), Inches(12), Inches(5.6), size=18)

    # 12. Next steps
    sl = deck.slide('Next steps', notes='Plan toward the CVPR submission (deadline November 16).')
    deck.bullets(sl, [
        'Robustness check on a second dataset: Amazon Shopping Queries Image Dataset (human-labeled exact matches)',
        'Optional: text-only vs image condition on a subset of models',
        'Optional: forced-thinking condition for Claude (explicit reasoning budget) to strengthen the effort manipulation',
        'Paper draft following CVPR guidelines',
    ], Inches(0.7), Inches(1.4), Inches(12), Inches(5.6), size=19)

    # Backup: position bias
    sl = deck.slide('Backup: position bias', notes=(
        'The preliminary run showed a primacy effect; with same-category distractors in the full study it is gone.'))
    sl.shapes.add_picture(figs['position'], Inches(0.4), Inches(1.3), width=Inches(12.5))
    deck.bullets(sl, [
        f'Position 1 is chosen {100 * s["pos_share"].loc[1]:.1f}% of the time (10% if position did not matter); '
        f'success is {100 * s["succ_front"]:.0f}% vs {100 * s["succ_back"]:.0f}% with the correct product in '
        'positions 1-5 vs 6-10',
        'The correct product\'s position was randomized and is identical across all conditions, so position cannot '
        'bias comparisons between models or effort levels',
    ], Inches(0.6), Inches(6.0), Inches(12.2), Inches(1.3), size=15)

    deck.prs.save(out_path)


def main():
    os.makedirs(FIG_DIR, exist_ok=True)
    d = load()
    s = compute_stats(d, np.random.default_rng(0))
    figs = {name: join(FIG_DIR, f'{name}.png') for name in
            ['effort', 'effects', 'calibration', 'right_wrong', 'concentration', 'position']}
    fig_effort(s, figs['effort'])
    fig_effects(s, figs['effects'])
    fig_calibration(s, figs['calibration'])
    fig_right_wrong(s, figs['right_wrong'])
    fig_concentration(s, figs['concentration'])
    fig_position(s, figs['position'])
    out_path = join(OUT_DIR, 'webshop_full_study_results.pptx')
    build(s, figs, out_path)
    print(f'wrote {out_path}')
    print('concentration: >=6 wrong obs %d vs exp %.1f | all right obs %d vs exp %.0f | modal wrong share %.2f vs random %.2f (%d tasks)'
          % (s['conc_obs_maj'], s['conc_exp_maj'], s['conc_obs_zero'], s['conc_exp_zero'], s['conc_modal'],
             s['conc_modal_random'], s['conc_n_multi']))
    print('position: front %.3f back %.3f | pos1 share %.3f' % (s['succ_front'], s['succ_back'], s['pos_share'].loc[1]))


if __name__ == '__main__':
    main()
